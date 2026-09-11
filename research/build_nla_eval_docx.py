#!/usr/bin/env python3
"""Build CoNaIISI student paper: NLA Eval."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

OUT = Path(__file__).resolve().parent / "conaiisi_2026_nla_eval.docx"
TNR = "Times New Roman"


def _rfonts(run, name: str = TNR) -> None:
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def set_run(run, size: int, *, bold=False, italic=False) -> None:
    run.font.name = TNR
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    _rfonts(run)


def add_p(
    doc,
    text: str,
    size: int,
    *,
    bold=False,
    italic=False,
    align="left",
    space_after=6,
    space_before=0,
    justify=False,
):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.0
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run(run, size, bold=bold, italic=italic)
    return p


def heading(doc, text: str, size: int = 12) -> None:
    add_p(doc, text, size, bold=True, space_before=10, space_after=6)


def body(doc, text: str, size: int = 12) -> None:
    add_p(doc, text, size, justify=True, space_after=6)


def set_a4(section) -> None:
    section.page_width = Twips(11907)
    section.page_height = Twips(16840)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Twips(1418)
    section.bottom_margin = Twips(1418)
    section.left_margin = Twips(1418)
    section.right_margin = Twips(1418)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)


def _strip_cols_and_page_num(sect_pr) -> None:
    for child in list(sect_pr):
        if child.tag in (qn("w:cols"), qn("w:pgNumType")):
            sect_pr.remove(child)


def set_title_block_section(section) -> None:
    set_a4(section)
    _strip_cols_and_page_num(section._sectPr)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), "none")
    section._sectPr.append(pg_num)


def set_two_equal_columns(section) -> None:
    set_a4(section)
    sect_pr = section._sectPr
    _strip_cols_and_page_num(sect_pr)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:equalWidth"), "0")
    cols.set(qn("w:num"), "2")
    col0 = OxmlElement("w:col")
    col0.set(qn("w:space"), "567")
    col0.set(qn("w:w"), "4252")
    col1 = OxmlElement("w:col")
    col1.set(qn("w:space"), "0")
    col1.set(qn("w:w"), "4252")
    cols.append(col0)
    cols.append(col1)
    sect_pr.append(cols)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), "none")
    sect_pr.append(pg_num)


def shade_header_cells(row) -> None:
    for cell in row.cells:
        tc = cell._tc
        tc_pr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8E8E8")
        shd.set(qn("w:val"), "clear")
        tc_pr.append(shd)


def fig_caption(doc, text: str) -> None:
    add_p(doc, text, 10, italic=True, align="center", space_before=8, space_after=8)


def add_table(doc, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run(run, 9, bold=True)
    shade_header_cells(table.rows[0])
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run(run, 8)


def main() -> None:
    doc = Document()
    set_title_block_section(doc.sections[0])

    add_p(
        doc,
        "Un marco de evaluación para Natural Language Autoencoders",
        16,
        bold=True,
        align="center",
        space_after=2,
    )
    add_p(
        doc,
        "De la inspección puntual al experimento reproducible",
        16,
        bold=True,
        align="center",
        space_after=12,
    )
    add_p(doc, " ", 14, bold=True, align="center", space_after=0)
    add_p(
        doc,
        "Universidad Tecnológica Nacional, [Unidad académica a completar]",
        12,
        bold=True,
        italic=True,
        align="center",
        space_after=6,
    )
    add_p(
        doc,
        "Nota. La versión para evaluación debe ser ciega. La sección Autores se deja en blanco y se completa solo en camera ready si el trabajo es aceptado.",
        10,
        italic=True,
        justify=True,
        space_after=10,
    )

    section1 = doc.add_section(WD_SECTION.CONTINUOUS)
    set_two_equal_columns(section1)

    heading(doc, "Abstract", 10)
    add_p(
        doc,
        (
            "El presente trabajo busca un primer acercamiento a la evaluación "
            "sistemática de verbalizaciones de activación producidas por Natural "
            "Language Autoencoders (NLA). Un NLA comprime el vector de activación "
            "de un modelo de lenguaje en un texto breve y reconstruye ese vector a "
            "partir del texto. Ese texto, llamado Activation Verbalizer (AV), es un "
            "señalizador interpretable, no una lectura literal de las creencias del "
            "modelo. Las plataformas actuales, en particular Neuronpedia, permiten "
            "inspeccionar un chat y un token a la vez. No ofrecen un ciclo de "
            "experimento comparable al que LangSmith formalizó para las salidas de "
            "aplicaciones con inteligencia artificial. Anthropic ya aplica un LLM "
            "como juez sobre AVs en un estudio de awareness de evaluación, pero ese "
            "juez vive dentro de un paper, no como herramienta reutilizable sobre NLA "
            "públicos. Se propone NLA Eval, un marco y un prototipo que orquesta "
            "datasets de prompts, una política de tokens, llamadas a la API de "
            "Neuronpedia, evaluadores configurables (léxico y LLM-as-judge) y una "
            "vista de comparación de corridas. El objeto evaluado es el AV en una "
            "posición de token, no la respuesta pública del modelo. La validación "
            "empírica a escala y el acuerdo humano-juez sobre un corpus etiquetado "
            "se plantean como trabajo inmediato."
        ),
        10,
        italic=True,
        justify=True,
        space_after=8,
    )

    heading(doc, "Palabras Clave", 10)
    add_p(
        doc,
        "Natural Language Autoencoders, interpretabilidad, LLM-as-judge, Neuronpedia, evaluación de modelos de lenguaje",
        10,
        space_after=12,
    )

    heading(doc, "1. Introducción")
    body(
        doc,
        "Los modelos de lenguaje de gran escala (LLM) se evalúan, en la práctica "
        "industrial, sobre lo que dicen. Suites de prompts, métricas y jueces "
        "automáticos comparan respuestas, costos y tasas de error. Ese enfoque "
        "alcanza cuando el objeto de interés es la salida. Quiebra cuando el "
        "fenómeno vive en las activaciones internas y no aparece en el texto que el "
        "usuario lee.",
    )
    body(
        doc,
        "Los Natural Language Autoencoders, introducidos por Fraser-Taliente y "
        "colaboradores en Transformer Circuits [1], atacan exactamente ese hueco. Un "
        "NLA es un autoencoder cuyo cuello de botella es texto en lenguaje natural. "
        "El Activation Verbalizer produce una verbalización. El Activation "
        "Reconstructor intenta recuperar el vector. El error de reconstrucción "
        "(MSE o RMSE) indica cuán leíble es esa verbalización, no si una frase "
        "aislada es verdadera.",
    )
    body(
        doc,
        "Neuronpedia, en colaboración con Anthropic, hostea checkpoints públicos "
        "(por ejemplo Llama 3.3 70B con kitft-l53 y Gemma 3 27B con kitft-l41) y "
        "una API de completion y explain [2]. El producto es un microscopio. Un "
        "prompt, un modelo, una pestaña. Quien investiga un tema (un prior de foro, "
        "un riesgo, una awareness de evaluación) termina exportando capturas, "
        "eligiendo a mano el último token del usuario y el primero del asistente, "
        "y pegando verbalizaciones en otro chat para que un LLM encuentre el "
        "patrón. Ese flujo no es un experimento. Es artesanía.",
    )
    body(
        doc,
        "En paralelo, el testing de aplicaciones no deterministas con IA ya tiene un "
        "marco de producto. LangSmith trata datasets, corridas, evaluadores (código "
        "o LLM-as-judge) y gráficos de comparación entre versiones de un flujo [3]. "
        "El objeto sigue siendo la salida de la aplicación.",
    )
    body(
        doc,
        "El presente trabajo propone el análogo para NLA. No se inventa el "
        "LLM-as-judge sobre verbalizaciones. Anthropic lo usó para marcar awareness "
        "de evaluación no verbalizada, con un grader sobre 50 tokens de respuesta "
        "y un 97 % de acuerdo con los autores en 186 ítems [1]. Se cambia el objeto "
        "evaluado y se formaliza el ciclo que hoy falta en las herramientas abiertas. "
        "El resto del artículo se organiza así. La Sección 2 formula las "
        "motivaciones. La Sección 3 resume el modelo conceptual. La Sección 4 "
        "describe la arquitectura, que es el aporte central. La Sección 5 describe "
        "el prototipo. La Sección 6 diseña la validación. La Sección 7 discute "
        "limitaciones. La Sección 8 concluye.",
    )

    heading(doc, "2. Motivaciones")
    body(
        doc,
        "Tres tensiones justifican el marco. Primero, el no determinismo. Un test "
        "clásico de API afirma que, dado un input, el output es uno. Un NLA, como "
        "el LLM que explica, no se presta a esa igualdad. Hace falta una política "
        "fija de tokens, un juez con rúbrica, y tasas a lo largo de un dataset, no "
        "una captura. Segundo, la escala. El límite de explain de Neuronpedia (del "
        "orden de 16 posiciones por pedido y un tope horario) convierte cada "
        "comparación Llama frente a Gemma en un trabajo de scripts y JSON. Sin un "
        "orquestador, el investigador elige el token interesante en cada corrida y "
        "destruye la comparabilidad. Tercero, el vacío de producto. LangSmith no "
        "conoce capas, residuos ni MSE. Neuronpedia no conoce datasets, claves de "
        "feedback ni gráficos de hit rate. EasyNLA incluye jueces de unicidad o "
        "coherencia mientras se entrena un NLA [1]. Ese es otro trabajo. No puntúa "
        "hipótesis experimentales sobre AVs ya desplegados.",
    )
    body(
        doc,
        "El motivo, en una frase, es este. Fomentar investigación sobre NLA, no el "
        "entrenamiento de NLA nuevos, sistematizando lo que hoy se hace a mano con "
        "capturas y un chat.",
    )

    heading(doc, "3. Modelo conceptual")
    body(
        doc,
        "El marco distingue cuatro objetos, en analogía deliberada con LangSmith. "
        "La Tabla 1 resume el mapeo.",
    )
    fig_caption(
        doc,
        "Tabla 1. Mapeo de objetos entre LangSmith y NLA Eval.",
    )
    add_table(
        doc,
        ["LangSmith", "NLA Eval"],
        [
            ("Dataset de prompts", "Igual. Prompt y, opcional, referencia"),
            ("Target = aplicación", "completion + explain bajo política de tokens"),
            ("Output de texto", "Artefacto completion, AVs, mse, token"),
            ("Claves de feedback", "Boolean, score o categoría"),
            ("Comparar experimentos", "Dos corridas sobre el mismo dataset"),
        ],
    )
    body(
        doc,
        "Una política de tokens, por defecto, replica el gesto que ya se usa en el "
        "laboratorio. Último token de contenido del mensaje de usuario. Primer token "
        "de la respuesta del asistente. El investigador puede elegir una, la otra, "
        "o ambas. El MSE actúa como compuerta. Un AV con error alto no se lee como "
        "evidencia. Se reporta como verbalización poco reconstruible.",
    )
    body(
        doc,
        "El juez no puntúa la creencia del modelo. Produce una señal interpretable. "
        "Si la rúbrica pide tema de foro o comunidad, un true no prueba que el "
        "preentrenamiento sea Reddit. Apoya una hipótesis que el output solo no "
        "mostraba, a validar con más tests [1].",
    )
    heading(doc, "3.1 Tipos de evaluador", 12)
    body(
        doc,
        "Léxico o código. Subcadena (reddit, forum). Barato y honesto. No captura "
        "tema. LLM-as-judge, sin referencia. Rúbrica Mustache sobre el AV. Boolean, "
        "score o categoría, más un campo opcional de razonamiento. Omisión, "
        "específico de NLA. El tema está en el AV y no en la completion. Es el "
        "análogo del awareness verbalizado frente al medido por NLA en [1]. "
        "LangSmith no tiene este tipo como objeto de primer orden. Un evaluador "
        "pairwise (cuál de dos AVs se prefiere) queda especificado en el diseño y no "
        "es el núcleo del prototipo actual.",
    )

    heading(doc, "4. Arquitectura")
    body(
        doc,
        "NLA Eval no entrena NLA. Orquesta APIs. La Figura 1 describe el ciclo de "
        "una fila. El dataset aporta el example (input). El run produce la completion "
        "y las verbalizaciones en las posiciones elegidas. El evaluador consume "
        "input, AV de entrada, AV de salida y, si hace falta, la completion, y calcula "
        "métricas.",
    )
    fig_caption(
        doc,
        "Figura 1. El evaluador compara el ejemplo del dataset con las verbalizaciones NLA del token de entrada y del token de salida, no solo input contra output de la aplicación.",
    )
    add_table(
        doc,
        ["Bloque", "Rol"],
        [
            ("Example", "Input = prompt del dataset"),
            ("Run", "Modelo + NLA. Output = completion"),
            ("NLA usuario", "Verbalización del último token de contenido"),
            ("NLA asistente", "Verbalización del primer token de respuesta"),
            ("Evaluator", "Métricas sobre AVs (y omisión vs completion)"),
        ],
    )
    body(
        doc,
        "La Figura 2 muestra el uso que motiva el marco. Dos configuraciones (v1 y "
        "v2) recorren el mismo dataset. El panel de evaluadores resume tasas (por "
        "ejemplo hit rate de un juez booleano) para decidir si un modelo más barato "
        "o una política de tokens distinta cambia el señalizador interno.",
    )
    fig_caption(
        doc,
        "Figura 2. Varias versiones de aplicación NLA (modelo, capa, política de tokens) sobre un único dataset. Los evaluadores agregan métricas por experimento.",
    )
    add_table(
        doc,
        ["Versión", "Ejemplo"],
        [
            ("v1", "Llama 3.3 70B Instruct, kitft-l53, último token de usuario"),
            ("v2", "Gemma 3 27B Instruct, kitft-l41, misma política"),
            ("Dataset", "Misma lista de prompts para ambas corridas"),
            ("Evaluadores", "Hit rate, score medio, categoría, MSE medio"),
        ],
    )

    heading(doc, "4.1 Interfaz", 12)
    body(
        doc,
        "La interfaz es un laboratorio web. Home de onboarding. Datasets. "
        "Evaluadores con configuración de feedback alineada a LangSmith (boolean, "
        "score, categórica, include reasoning). Corrida de experimento (fuente NLA, "
        "política, jueces). Tabla de filas con celdas coloreadas (cumple o no "
        "cumple, score, categoría) y gráficos de comparación. Las claves de OpenAI y "
        "Neuronpedia viajan en la sesión del navegador hacia un proxy. No se "
        "persisten en el cliente como secretos de producto.",
    )
    heading(doc, "4.2 Orquestador de experimento", 12)
    body(
        doc,
        "El orquestador es el análogo del kernel de ARIA [4], con un trabajo más "
        "estrecho. No clasifica requerimientos. Secuencia, por cada ejemplo del "
        "dataset, tres pasos. Primero, POST a la API de completion de Neuronpedia, "
        "con modelo e identificador de NLA (kitft-l53, kitft-l41). Segundo, "
        "selección de posiciones según la política y POST de explain. Tercero, para "
        "cada evaluador, un LLM u operación de código sobre el artefacto. El esquema "
        "JSON de salida lo dictan las claves de feedback.",
    )
    body(
        doc,
        "El orquestador persiste el experimento (filas, scores, comentarios, MSE) "
        "para que un GET posterior, en otro proceso, no pierda la corrida. En "
        "servidor local el almacén es un JSON. En despliegue el almacén debe ser "
        "compartido (por ejemplo Postgres). El filesystem efímero de una función "
        "serverless no es un laboratorio.",
    )
    heading(doc, "4.3 Capa NLA", 12)
    body(
        doc,
        "Esta capa es un cliente. No reimplementa el autoencoder. Encapsula rate "
        "limits y el recorte de posiciones. Si explain falla, la fila queda con error "
        "y no se inventa un AV. El prototipo usa Llama 3.3 70B Instruct más Gemma 3 "
        "27B como contraste, no como par equivalente de claim científico.",
    )
    heading(doc, "4.4 Capa juez", 12)
    body(
        doc,
        "Los prompts del juez son plantillas Mustache. El mapeo fija qué variable "
        "es el AV agregado, el AV de último token de usuario, el de primer token de "
        "asistente, el prompt, la completion, el token, el MSE o una referencia "
        "humana. El formato de respuesta (boolean, score con min y max, categorías "
        "con descripción) estructura el JSON. El razonamiento es opcional. El juez "
        "se aplica al AV, no a la respuesta pública, salvo que la rúbrica lo pida "
        "(omisión).",
    )
    heading(doc, "4.5 Capa de comparación", 12)
    body(
        doc,
        "La tabla de un run copia el gesto de LangSmith. Columnas de input, "
        "referencia, output (el AV) y una columna por clave de feedback. Boolean se "
        "muestra como 1.00 o 0.00 con color. Score se colorea entre mínimo y máximo. "
        "Categórica muestra el nombre. El encabezado informa el promedio. La "
        "comparación de dos experimentos reusa las mismas claves sobre el mismo "
        "dataset. Un evaluador pairwise (Figura 3) queda como extensión. Dos AVs del "
        "mismo example, un juez de preferencia, preferred en {0, 1}.",
    )
    fig_caption(
        doc,
        "Figura 3. Evaluación pairwise (diseñada, no núcleo del prototipo). El juez elige entre dos verbalizaciones del mismo example.",
    )
    add_table(
        doc,
        ["Paso", "Descripción"],
        [
            ("v1 y v2", "Dos corridas sobre el mismo dataset"),
            ("Par", "AVs del mismo example"),
            ("Juez pairwise", "preferred = 0 o 1, más razonamiento opcional"),
        ],
    )

    heading(doc, "5. Prototipo")
    body(
        doc,
        "Se materializó un prototipo web (Next.js) que implementa el ciclo de las "
        "Secciones 4.1 a 4.5 contra la API real de Neuronpedia y un juez OpenAI. El "
        "prototipo no es la contribución teórica. Es la prueba de que el marco cabe "
        "en un frontend más dos llamadas, como se anticipaba. Dataset semilla de "
        "prompts de consejo y how-to, juez de tema foro o comunidad, política de "
        "último token de usuario, tabla de filas y gráficos de hit rate y MSE medio.",
    )
    body(
        doc,
        "No se reportan aquí tasas como hallazgo. Un prototipo de pocos días no "
        "sustituye un protocolo de validación. Sí muestra que el cuello de botella "
        "deja de ser cómo pego diez capturas y pasa a ser qué rúbrica y qué "
        "política de tokens son honestas.",
    )

    heading(doc, "6. Diseño de validación")
    body(
        doc,
        "La validación se plantea en dos niveles, como en [4]. A nivel de señal, un "
        "conjunto fijo de prompts (familia consejo, how-to, enciclopedia, un ítem en "
        "español). Dos fuentes NLA (Llama 70B kitft-l53 y Gemma 27B kitft-l41). Una "
        "política de tokens congelada. Un juez booleano de tema de foro y un "
        "evaluador léxico. Métricas. Hit rate del juez, tasa léxica, tasa de omisión "
        "(AV sí, completion no), MSE medio. Acuerdo entre el juez y etiquetado "
        "humano sobre una muestra de AVs, en el espíritu del 97 % reportado en [1], "
        "sin copiar ese número.",
    )
    body(
        doc,
        "A nivel de herramienta. Tiempo para reproducir una comparación que hoy exige "
        "pestañas. Capacidad de un tercero de definir un juez distinto (por ejemplo "
        "peligro o evaluación) sin tocar Neuronpedia. Usabilidad del laboratorio. Los "
        "resultados de esa validación son el trabajo que este artículo deja "
        "explícitamente abierto, no un resultado adelantado.",
    )

    heading(doc, "7. Limitaciones")
    body(
        doc,
        "El marco hereda los límites del instrumento. El NLA se entrena en una capa "
        "fija. El AV es un señalizador. El juez puede sesgarse hacia lexemas o hacia "
        "el propio prompt. Neuronpedia impone cupos. El prototipo no cubre pairwise "
        "ni jueces de código más allá del léxico. No hay NLA públicos sobre modelos "
        "base, así que no se replica el eje base frente a SFT de otros papers de "
        "sesgo cultural [5]. Nada de esto autoriza a leer un AV como pensamiento.",
    )

    heading(doc, "8. Conclusiones y trabajos futuros")
    body(
        doc,
        "Se presentó un marco para evaluar verbalizaciones NLA con el mismo ciclo "
        "que LangSmith ofrece para salidas de aplicaciones con IA. El aporte no es "
        "un método nuevo de interpretabilidad. Es convertir la inspección puntual de "
        "Neuronpedia y el grader ocasional de Anthropic en un laboratorio con "
        "dataset, política de tokens, evaluadores y comparación de corridas. El "
        "prototipo muestra que la implementación es un orquestador sobre APIs "
        "existentes. La validación empírica, el acuerdo humano-juez y el evaluador "
        "pairwise son los pasos siguientes.",
    )
    body(
        doc,
        "Neuronpedia fomentó el acceso a NLA. NLA Eval apunta a fomentar "
        "investigaciones que los usen, con una forma determinística de probar "
        "hipótesis sobre temas en AVs.",
    )

    heading(doc, "Agradecimientos", 10)
    add_p(
        doc,
        "[Nombres y apellidos de los docentes tutores. Completar en camera ready.]",
        10,
        space_after=8,
    )

    heading(doc, "Referencias", 10)
    refs = [
        "[1] K. Fraser-Taliente et al., “Natural Language Autoencoders,” Transformer Circuits, may. 2026. https://transformer-circuits.pub/2026/nla/",
        "[2] Neuronpedia, “Natural Language Autoencoders demo and API.” https://www.neuronpedia.org/nla",
        "[3] LangChain, “LangSmith: evaluation and experiment tracking for LLM applications.” https://docs.smith.langchain.com/",
        "[4] P. Kozak Riehme et al., “ARIA: Un Enfoque Inteligente para la Gestión y Mejora de Requerimientos en Entornos Ágiles,” CoNaIISI, 2026.",
        "[5] J. Fernandez de Landa et al., “Why are all LLMs Obsessed with Japanese Culture? On the Hidden Cultural and Regional Biases of LLMs,” Findings of EMNLP, 2026. arXiv:2604.21751.",
        "[6] Anthropic, “Natural Language Autoencoders,” 2026. https://www.anthropic.com/research/natural-language-autoencoders",
        "[7] kitft, “natural_language_autoencoders,” GitHub. https://github.com/kitft/natural_language_autoencoders",
    ]
    for t in refs:
        add_p(doc, t, 10, justify=True, space_after=4)

    heading(doc, "Datos de Contacto", 10)
    add_p(
        doc,
        "Completar en camera ready. Nombre y Apellido. Universidad Tecnológica Nacional, [unidad académica]. Dirección postal. E-mail.",
        10,
        italic=True,
        justify=True,
    )

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
