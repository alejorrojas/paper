#!/usr/bin/env python3
"""Build CoNaIISI student paper matching Formato_Estudiantes."""

from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from pathlib import Path

OUT = Path(__file__).resolve().parent / "conaiisi_2026_reddit_nla.docx"
TNR = "Times New Roman"


def _rfonts(run, name: str = TNR) -> None:
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def set_run(run, size: int, *, bold=False, italic=False) -> None:
    run.font.name = TNR
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    _rfonts(run)


def add_p(doc, text: str, size: int, *, bold=False, italic=False, align="left", space_after=6, space_before=0, justify=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.0
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run(run, size, bold=bold, italic=italic)
    return p


def heading(doc, text: str, size: int = 12) -> None:
    add_p(doc, text, size, bold=True, space_before=10, space_after=6)


def body(doc, text: str, size: int = 12) -> None:
    add_p(doc, text, size, justify=True, space_after=6)


def set_a4(section) -> None:
    # Formato_Estudiantes (Google Doc export): A4 11907×16840 twips, 2.5 cm = 1418.
    section.page_width = Twips(11907)
    section.page_height = Twips(16840)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Twips(1418)
    section.bottom_margin = Twips(1418)
    section.left_margin = Twips(1418)
    section.right_margin = Twips(1418)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)


def _strip_cols_and_page_num(sect_pr) -> None:
    for child in list(sect_pr):
        if child.tag in (qn("w:cols"), qn("w:pgNumType")):
            sect_pr.remove(child)


def set_title_block_section(section) -> None:
    """Single column, like ANEXO II title block in Formato_Estudiantes."""
    set_a4(section)
    _strip_cols_and_page_num(section._sectPr)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), "none")
    section._sectPr.append(pg_num)


def set_two_equal_columns(section) -> None:
    """Two equal columns, 1 cm gap — copied from Formato_Estudiantes sectPr."""
    set_a4(section)
    sect_pr = section._sectPr
    _strip_cols_and_page_num(sect_pr)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:equalWidth"), "0")
    cols.set(qn("w:num"), "2")
    col0 = OxmlElement("w:col")
    col0.set(qn("w:space"), "567")
    col0.set(qn("w:w"), "4252")
    col1 = OxmlElement("w:col")
    col1.set(qn("w:space"), "0")
    col1.set(qn("w:w"), "4252")
    cols.append(col0)
    cols.append(col1)
    sect_pr.append(cols)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), "none")
    sect_pr.append(pg_num)


def shade_header_cells(row) -> None:
    for cell in row.cells:
        tc = cell._tc
        tc_pr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8E8E8")
        shd.set(qn("w:val"), "clear")
        tc_pr.append(shd)


def main() -> None:
    doc = Document()
    section0 = doc.sections[0]
    set_title_block_section(section0)

    add_p(
        doc,
        "¿Reddit como fuente de verdad para la IA?",
        16,
        bold=True,
        align="center",
        space_after=2,
    )
    add_p(
        doc,
        "Sobre los priors de foro ocultos en verbalizaciones NLA",
        16,
        bold=True,
        align="center",
        space_after=12,
    )
    add_p(doc, " ", 14, bold=True, align="center", space_after=0)
    add_p(
        doc,
        "Universidad Tecnológica Nacional, [Unidad académica a completar]",
        12,
        bold=True,
        italic=True,
        align="center",
        space_after=6,
    )
    add_p(
        doc,
        "Nota. La versión para evaluación debe ser ciega. La sección Autores se deja en blanco y se completa solo en camera ready si el trabajo es aceptado.",
        10,
        italic=True,
        justify=True,
        space_after=10,
    )

    # Official template: 1-col title block, then continuous 2-col from Abstract.
    section1 = doc.add_section(WD_SECTION.CONTINUOUS)
    set_two_equal_columns(section1)

    heading(doc, "Abstract", 10)
    add_p(
        doc,
        (
            "Los modelos de lenguaje tienen límites de cobertura y, en algunos casos, "
            "defaults de fuente. Aunque hay estudios sobre sesgos culturales y sobre la "
            "calidad de las respuestas, ninguno ha investigado de forma específica qué "
            "comunidad eligen ante preguntas genéricas de consejo. En este trabajo se "
            "propone un conjunto de pedidos abiertos con la comunidad o el sitio "
            "enmascarados. Se evalúa Llama 3.3 70B Instruct pidiéndole que responda y "
            "elija ese ancla. En una segunda batería se aplica un Natural Language "
            "Autoencoder (NLA) al último token del usuario. Los resultados muestran una "
            "tendencia clara hacia Reddit cuando hay que elegir dónde se habla del tema. "
            "Si el hueco es un sitio web, el modelo elige una tienda. Además, en pedidos "
            "de consejo sin ancla el NLA nombra Reddit y el completion no lo cita, salvo "
            "una lista de sitios. En salud o finanzas la verbalización se alinea con un "
            "artículo o una FAQ. Por último, se replica el protocolo en Gemma 3 27B. Ese "
            "modelo cita subreddits en el slot de comunidad y no muestra el lexema Reddit "
            "en el NLA."
        ),
        10,
        italic=True,
        justify=True,
        space_after=8,
    )

    heading(doc, "Palabras Clave", 10)
    add_p(
        doc,
        "Inteligencia Artificial, Natural Language Autoencoders, sesgo de default, Reddit, interpretabilidad, Llama 3.3 70B",
        10,
        space_after=12,
    )

    heading(doc, "Introducción")
    body(
        doc,
        "Se espera que, ante la misma pregunta de consejo, las personas anclen la "
        "respuesta a comunidades distintas. El comportamiento de los modelos de "
        "lenguaje ante ese tipo de pedido abierto sigue poco descrito. Hay trabajo que "
        "evalúa sesgos contra un conjunto de respuestas correctas o contra respuestas "
        "humanas. En este trabajo se propone un marco para medir priors de comunidad. "
        "Los prompts son abiertos e ambiguos. La información de foro o de sitio está "
        "enmascarada. El modelo debe responder y anclar esa respuesta a un lugar donde "
        "se habla del tema. Al forzar el ancla, el marco saca a la superficie priors "
        "que una respuesta sin ancla no muestra. En una segunda capa, un Natural "
        "Language Autoencoder (NLA) [1] verbaliza activaciones en el último token del "
        "usuario, incluso cuando el texto generado no cita la plataforma.",
    )
    body(
        doc,
        "Las pruebas se corrieron con el NLA público de Neuronpedia [4], [5] sobre "
        "Llama 3.3 70B Instruct, capa 53 (kitft l53). No hay NLA público para Llama 4. "
        "Un contraste con Gemma 3 27B Instruct, NLA capa 41, evita atribuir a todos "
        "los modelos lo que puede ser propio de uno.",
    )
    body(
        doc,
        "Se abordan las siguientes preguntas. Qué comunidad elige el modelo cuando "
        "debe elegir dónde se habla del tema, sin que el prompt nombre Reddit. Si "
        "ocurre lo mismo cuando el hueco es un sitio web. Si, sin ese ancla, las "
        "verbalizaciones del último token del usuario nombran Reddit aunque la "
        "respuesta no lo cite. Si ese rastro interno se replica en Gemma. El resto "
        "del trabajo se organiza en marco teórico, metodología, resultados en Llama, "
        "contraste con Gemma, trabajos relacionados, conclusión y trabajos futuros.",
    )

    heading(doc, "Marco teórico")
    heading(doc, "2.1 Priors de comunidad en preguntas abiertas", 12)
    body(
        doc,
        "Una pregunta de consejo subespecificada deja un hueco. El modelo tiene que "
        "completar de dónde sale la recomendación. Si se le pide explícitamente que "
        "elija la comunidad o el sitio, ese hueco deja de ser implícito. No se "
        "afirma que Reddit entrenó al modelo. Se mide qué nombre llena el hueco y "
        "qué temas aparecen en el NLA. Se llama default visible al que sale en el "
        "completion, por ejemplo citar un subreddit. Se llama prior oculto al lexema "
        "Reddit o foro en la verbalización cuando el completion es una guía sin citar "
        "la plataforma. El segundo es lo que el output solo no muestra.",
    )
    heading(doc, "2.2 Natural Language Autoencoders", 12)
    body(
        doc,
        "Un NLA es un autoencoder cuyo cuello de botella es texto en lenguaje natural "
        "[1]. Dada una activación h_l del residual stream en la capa l, el Activation "
        "Verbalizer (AV) genera un texto z. El Activation Reconstructor (AR) intenta "
        "reconstruir h_l desde z. El costo, de cientos de tokens por activación, obliga "
        "a sondeos dispersos. El criterio de lectura privilegia temas recurrentes y no "
        "frases aisladas. Las métricas RMSE y FVE estiman reconstrucción y no la "
        "verdad semántica de cada oración [1]. En atención causal, la AV en el token "
        "t describe el prefijo x_1 a x_t. El modelo no ve si el usuario iba a seguir "
        "tipeando.",
    )

    heading(doc, "Metodología")
    body(
        doc,
        "Se corrieron pedidos de géneros distintos en Free Chat y por API, siempre con "
        "el NLA público y sin nombrar Reddit. Sobre esa base se fijaron dos baterías "
        "para cuantificar el patrón y un contraste en Gemma. No se entrena un NLA. El "
        "completion usa temperatura 0,4 y el explain 0,7. Las once capturas de Llama se "
        "tomaron en Free Chat de Neuronpedia y el resto por API. El modelo principal es "
        "Llama 3.3 70B Instruct, NLA capa 53. El contraste es Gemma 3 27B Instruct, NLA "
        "capa 41, con los mismos prompts.",
    )
    heading(doc, "Experimento A. Comunidad o sitio enmascarados", 12)
    body(
        doc,
        "Se usaron cuatro prompts en inglés, sin la palabra Reddit. Los dos primeros "
        "tratan laptop y amigos con el placeholder according to community y la "
        "instrucción Choose yourself where people talk about this. El tercero es "
        "laptop sin slot. El cuarto es laptop con on website y Choose yourself the "
        "website. Se registra si el completion nombra Reddit o un subreddit.",
    )
    heading(doc, "Experimento B. Sin ancla, último token del usuario", 12)
    body(
        doc,
        "Se usaron once prompts de uso cotidiano, entre ellos resumir, calorías, "
        "fuente de verdad en una palabra, laptop, arroz, lista de sitios, Give me "
        "ideas for, correo profesional no rígido, sueño en turnos de noche, Roth "
        "frente a Traditional IRA, y equity en una oferta de startup en español. Se "
        "explica el último token de contenido del mensaje de usuario. El criterio "
        "binario pregunta si aparece la cadena Reddit, sin distinguir mayúsculas, en "
        "la AV y si aparece en el completion. Un mix adicional de ocho pedidos "
        "(invertir, amigos, Python, Fibonacci, mudanza, arroz, laptop y endpoint "
        "vulnerable) sondea los últimos cuatro tokens de usuario y el primero del "
        "assistant. Se reporta presencia de Reddit en alguna AV de esa ventana.",
    )
    heading(doc, "Qué no se mide", 12)
    body(
        doc,
        "No hay NLA sobre modelos base, es decir sin instruct. No se recorre toda la "
        "frase token a token. No se separa el aporte del pretrain respecto del verbalizer.",
    )

    heading(doc, "Resultados")
    heading(doc, "Default visible (Llama)", 12)
    body(
        doc,
        "En Llama, el slot de comunidad produce Reddit en el output, r/laptops y "
        "r/college para laptop, r/newcity y Reddit community para amigos. El slot "
        "de sitio no. Elige Best Buy. El laptop sin slot no cita Reddit y da una lista "
        "de marcas. Reddit es el default de comunidad y no de cualquier hueco "
        "semántico.",
    )
    heading(doc, "Prior oculto (Llama)", 12)
    body(
        doc,
        "La Tabla 1 resume el último token del usuario en las 11 capturas de Llama. "
        "El lexema Reddit aparece en el 73 % de las AV (8 de 11) y en el 9 % de los "
        "completions (1 de 11). La lista de diez sitios incluye reddit.com. En el mix "
        "de ocho pedidos, el 100 % (8 de 8) tiene Reddit en alguna AV de la cola del "
        "usuario y el 0 % (0 de 8) en el completion. No es un artefacto de "
        "ciberseguridad. Arroz, mudanza y Recommend a laptop for a student disparan el "
        "tema. Tampoco es solo inglés. Qué significa equity en una oferta de startup "
        "verbaliza Reddit y medios latinos.",
    )

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(4)
    r = cap.add_run(
        "Tabla 1. Llama 70B, capa 53, último token del usuario. Reddit en AV indica lexema en la verbalización."
    )
    set_run(r, 10, italic=True)

    table = doc.add_table(rows=12, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Prompt (gist)", "Token", "Reddit AV", "Reddit output"]
    data = [
        ("Summarize this for me", "me", "sí", "no"),
        ("Calories in a banana", "?", "no", "no"),
        ("Source of truth, one word", "word", "sí", "no (Wikipedia)"),
        ("Laptop for a student", "student", "sí", "no"),
        ("Cook white rice", "rice", "sí", "no"),
        ("List of 10 websites", "websites", "sí", "sí (ítem 6)"),
        ("Give me ideas for", "for", "sí", "no"),
        ("Email professional not stiff", "stiff", "sí", "no"),
        ("Night shift sleep", "shifts", "no", "no"),
        ("Roth vs Traditional IRA", "IRA", "no", "no"),
        ("Equity en startup (ES)", "startup", "sí", "no"),
    ]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run(run, 10, bold=True)
    shade_header_cells(table.rows[0])
    for i, row in enumerate(data, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run(run, 9)

    heading(doc, "Registro del prefijo", 12)
    body(
        doc,
        "Las tres celdas sin Reddit en Llama (calorías, sueño e IRA) se verbalizan como "
        "FAQ, CNN, Health.com o artículo de finanzas. No es solo salud. IRA no es "
        "medicina. La lectura que se propone es la siguiente. En el último token de una "
        "pregunta factual ya cerrada, el hilo del residual es encabezado de artículo. "
        "En un tutorial o consejo, el género del prefijo es hilo de foro. Eso es "
        "atención causal. En t el modelo no sabe si vendrán más tokens. Recorrer IRA "
        "token a token queda abierto.",
    )

    heading(doc, "Contraste. ¿Se replica en Gemma?")
    body(
        doc,
        "Las pruebas de verbalización se corrieron sobre Llama. Para no convertir "
        "un hallazgo de un modelo en un rasgo de toda la familia, se repitió el "
        "Experimento B en Gemma 3 27B Instruct, NLA capa 41, con los mismos once "
        "prompts y el último token del usuario, por API. El resultado fue 0 % (0 de 11) "
        "Reddit en AV y 0 % (0 de 11) en el completion. Una AV dice Chat o forum post "
        "sin nombrar Reddit.",
    )
    body(
        doc,
        "En el Experimento A, Gemma sí cita subreddits cuando el slot es comunidad "
        "(r/SuggestALaptop, r/AskReddit) y Best Buy cuando el slot es sitio. El default "
        "visible de comunidad, entonces, no es exclusivo de Llama. El lexema Reddit "
        "en el NLA, que es el resultado que el output no muestra, no se replica. "
        "Ese funcionamiento interno puede deberse a un sesgo de Llama o de su "
        "verbalizer kitft l53, y no a un prior universal de la IA. El contraste no "
        "prueba inocencia del verbalizer. Sí impide el slogan de que todos los LLM "
        "piensan en Reddit.",
    )

    heading(doc, "Trabajos Relacionados")
    body(
        doc,
        "Fraser-Taliente et al. [1] presentan los NLA, advierten confabulación de "
        "detalles y muestran casos en los que el AV añade señal que el output no "
        "verbaliza. En CoNaIISI 2024, Mansilla et al. [2] usan autoencoders clásicos "
        "para un primer acercamiento a una herramienta UEBA, con problema concreto, "
        "marco teórico didáctico y validación acotada. Un trabajo de cátedra sobre "
        "prejuicios e imparcialidades en IA [6] encuadra el tema ético con encuesta. "
        "Aquí el objeto es un default medible en prompts y verbalizaciones, y no la "
        "opinión pública. Requelme et al. [7] combinan LLM y grafos de conocimiento "
        "y documentan alucinaciones. Conviene leer las AV con el mismo recelo.",
    )
    body(
        doc,
        "Fernandez de Landa et al. [3] miden sesgos culturales y regionales con "
        "preguntas culturales abiertas y un placeholder de lugar. El modelo debe "
        "elegir la región. Reportan una tendencia hacia países como Japón, efectos "
        "de idioma, y señales de que el sesgo se ve más claro después del ajuste "
        "supervisado que en el pretrain. Su protocolo de pregunta abierta y ancla "
        "forzada es el precedente metodológico más cercano. El nuestro cambia el "
        "slot, comunidad o sitio y no país, y añade la capa NLA. No se afirma el "
        "mismo fenómeno cultural. Se cita el trabajo porque el marco de hueco "
        "enmascarado es el que se usa.",
    )

    heading(doc, "Conclusión y Trabajos Futuros")
    body(
        doc,
        "El AV puede confabular fuentes. No se demuestra origen en el corpus de "
        "pretrain. Las baterías cuantificadas son acotadas y conviene ampliarlas.",
    )
    body(
        doc,
        "En Llama 70B hay una tendencia a Reddit cuando se pide elegir comunidad. No "
        "ocurre si el hueco es un sitio. Sin ancla, el NLA de capa 53 nombra Reddit en "
        "el 73 % de once pedidos de consejo (8 de 11) y en el 100 % de otros ocho "
        "(8 de 8). El completion lo cita en el 9 % y en el 0 %. Gemma cita subreddits "
        "en el slot de comunidad y da 0 % en el NLA. El default visible de comunidad no "
        "es exclusivo de Llama. El lexema interno puede serlo.",
    )
    body(
        doc,
        "Los trabajos futuros apuntan a ampliar las baterías, recorrer preguntas "
        "factuales token a token, repetir seeds y, si aparece un NLA de Llama 4 u "
        "otros instruct, repetir el protocolo. También se busca separar verbalizer de "
        "modelo.",
    )

    heading(doc, "Agradecimientos", 10)
    add_p(
        doc,
        "[Nombres y apellidos de los docentes tutores. Completar en camera ready.]",
        10,
        space_after=8,
    )

    heading(doc, "Referencias", 10)
    refs = [
        "[1] K. Fraser-Taliente et al., “Natural Language Autoencoders Produce Unsupervised Explanations of LLM Activations”, Transformer Circuits Thread, 2026. https://transformer-circuits.pub/2026/nla/index.html",
        "[2] “Aplicación de Autoencoders en la Detección de Amenazas Internas: Estrategias y Resultados”, Memorias CoNaIISI 2024.",
        "[3] J. Fernandez de Landa, C. Perez-Almendros y J. Camacho-Collados, “Why are all LLMs Obsessed with Japanese Culture? On the Hidden Cultural and Regional Biases of LLMs”, EMNLP 2026 Findings. arXiv:2604.21751.",
        "[4] Neuronpedia, demo e API NLA. https://www.neuronpedia.org/nla",
        "[5] Anthropic, “Natural Language Autoencoders” (blog). https://www.anthropic.com/research/natural-language-autoencoders",
        "[6] I. Olmos et al., “Prejuicios e Imparcialidades en las Inteligencias Artificiales”, Memorias CoNaIISI 2024.",
        "[7] A. E. Requelme Ontiveros et al., “Aplicación de LLMs y KGs para la Interpretación de Bases de Datos Relacionales”, Memorias CoNaIISI 2024.",
    ]
    for t in refs:
        add_p(doc, t, 10, justify=True, space_after=4)

    heading(doc, "Datos de Contacto", 10)
    add_p(
        doc,
        "Completar en camera ready. Nombre y Apellido. Universidad Tecnológica Nacional, [unidad académica]. Dirección postal. E-mail.",
        10,
        italic=True,
        justify=True,
    )

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
